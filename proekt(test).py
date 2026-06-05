import tkinter as tk
import tkinter.ttk as ttk
import tkinter.messagebox
import tkinter.simpledialog
import tkinter.filedialog
import matplotlib.pyplot as plt
import numpy as np
import scipy.stats as sps
import scipy.interpolate as spi
import wordcloud as wc
import nltk
import pypdf
import docx
import statistics
import random
import collections
import os
import re
import copy

try:
    import openpyxl
    XL_OK = True
except ImportError:
    XL_OK = False

# Загрузка ресурсов NLTK
for res in ['tokenizers/punkt', 'corpora/stopwords']:
    try:
        nltk.data.find(res)
    except LookupError:
        nltk.download(res.split('/')[-1])

# -------------------------------------------------------------
# ГЛАВНОЕ ПРИЛОЖЕНИЕ
# -------------------------------------------------------------
class App:
    def __init__(self, master):
        self.master = master
        master.title("Обработка результатов эксперимента")
        master.geometry("1100x750")
        master.protocol("WM_DELETE_WINDOW", self._close)

        top = ttk.Frame(master)
        top.pack(pady=5, fill='x')

        self.btn_num = ttk.Button(top, text="📊 Числовые данные", command=self._show_num)
        self.btn_num.pack(side=tk.LEFT, padx=10)

        self.btn_txt = ttk.Button(top, text="📄 Текстовые данные", command=self._show_txt)
        self.btn_txt.pack(side=tk.LEFT, padx=10)

        self.btn_undo = ttk.Button(top, text="↶", command=self._undo)
        self.btn_undo.pack(side=tk.LEFT, padx=2)

        self.btn_redo = ttk.Button(top, text="↷", command=self._redo)
        self.btn_redo.pack(side=tk.LEFT, padx=2)

        self.btn_help = ttk.Button(top, text="❓ Справка", command=self._help)
        self.btn_help.pack(side=tk.RIGHT, padx=10)

        self.content = ttk.Frame(master)
        self.content.pack(fill='both', expand=True, padx=10, pady=10)

        self.num_frame = ttk.Frame(self.content)
        self.txt_frame = ttk.Frame(self.content)

        self.num_tab = NumTab(self.num_frame, self)
        self.txt_tab = TxtTab(self.txt_frame, self)

        self._show_num()
        master.bind('<Control-Key>', self._global_hotkey)

    def _global_hotkey(self, event):
        ks = event.keysym.lower()
        if ks in ('z', 'я'):
            self._undo()
        elif ks in ('y', 'н'):
            self._redo()

    def _show_num(self):
        self.txt_frame.pack_forget()
        self.num_frame.pack(fill='both', expand=True)
        self.btn_num.config(state='disabled')
        self.btn_txt.config(state='normal')

    def _show_txt(self):
        self.num_frame.pack_forget()
        self.txt_frame.pack(fill='both', expand=True)
        self.btn_txt.config(state='disabled')
        self.btn_num.config(state='normal')

    def _undo(self):
        if self.num_frame.winfo_ismapped():
            self.num_tab.undo()
        else:
            self.txt_tab.undo()

    def _redo(self):
        if self.num_frame.winfo_ismapped():
            self.num_tab.redo()
        else:
            self.txt_tab.redo()

    def _help(self):
        text = """🔹 КАК РАБОТАТЬ С ПРОГРАММОЙ 🔹

1. ЧИСЛОВЫЕ ДАННЫЕ:
   • Выберите режим ввода: "1 столбец" или "2 столбца" (кнопка переключения).
   • Для одномерных данных: введите число в поле и нажмите Enter.
   • Для двумерных данных (X,Y): введите X и Y в соответствующие поля и нажмите Enter.
   • Двойной клик по числу/паре – редактирование.
   • Кнопки: очистить, гистограмма, график, сортировка, фильтр, сброс фильтра.
   • Поиск, загрузка из файлов (TXT, CSV, Excel, Word), сохранение.
   • Загрузка двумерных данных также через "📂 Открыть (2 столбца)".
   • Математические методы – меню «📐 Мат. методы»:
        - Для двумерных (X,Y): корреляция Пирсона, линейная регрессия, полином 4-й степени,
          интерполяция (Лагранж, Ньютон, канонический, сплайны).
        - Для одномерных: t-тест, удаление выбросов, нормализация, стандартизация,
          доверительный интервал.
   • Отмена/повтор: Ctrl+Z, Ctrl+Y или кнопки ↶/↷ в главном окне.

2. ТЕКСТОВЫЕ ДАННЫЕ:
   • Загрузить TXT, CSV, PDF, DOCX или ввести вручную.
   • Поддерживаются русский и английский языки.
   • Кнопки: загрузить, сохранить, очистить, поиск, облако слов, удалить стоп-слова, стемминг.
   • Тональность анализируется автоматически (справа).
   • Отмена/повтор: Ctrl+Z, Ctrl+Y.

Горячие клавиши:
   Enter – добавить (числовой режим)
   Ctrl+Z – отменить
   Ctrl+Y – повторить
"""
        tk.messagebox.showinfo("Справка", text)

    def _close(self):
        try:
            plt.close('all')
        except:
            pass
        self.master.destroy()


# -------------------------------------------------------------
# ВКЛАДКА ЧИСЛОВЫХ ДАННЫХ (без кнопки "Добавить")
# -------------------------------------------------------------
class NumTab:
    def __init__(self, parent, app):
        self.parent = parent
        self.app = app
        self.data = []
        self.fdata = None
        self.f_active = False
        self.bivar = False
        self.x = []
        self.y = []
        self.two_column_mode = False
        self.hist = []
        self.hist_pos = -1
        self._save_hist()
        self._build()
        self.refresh_info()

    def _save_hist(self):
        self.hist = self.hist[:self.hist_pos+1]
        self.hist.append(copy.deepcopy(self.data))
        self.hist_pos += 1
        if len(self.hist) > 100:
            self.hist.pop(0)
            self.hist_pos -= 1

    def undo(self):
        if self.hist_pos > 0:
            self.hist_pos -= 1
            self.data = copy.deepcopy(self.hist[self.hist_pos])
            if self.f_active:
                self.unfilter()
            self._update_xy_from_data()
            self.show()
            self.refresh_info()

    def redo(self):
        if self.hist_pos < len(self.hist)-1:
            self.hist_pos += 1
            self.data = copy.deepcopy(self.hist[self.hist_pos])
            if self.f_active:
                self.unfilter()
            self._update_xy_from_data()
            self.show()
            self.refresh_info()

    def _update_xy_from_data(self):
        if self.bivar and self.data and isinstance(self.data[0], tuple):
            self.x = [p[0] for p in self.data]
            self.y = [p[1] for p in self.data]
        else:
            self.x = []
            self.y = []

    def _build(self):
        btn_cont = ttk.Frame(self.parent)
        btn_cont.pack(pady=10, fill='x')

        def btn(parent, text, cmd):
            b = ttk.Button(parent, text=text, command=cmd)
            b.pack(side=tk.LEFT, padx=2, pady=2)
            return b

        row1 = ttk.Frame(btn_cont)
        row1.pack(fill='x', pady=2)
        for (text, cmd) in [
            ("🗑️ Очистить", self.clear),
            ("📊 Гистограмма", self.hist_plot),
            ("📈 График", self.graph),
            ("🔽 Сортировать", self.sort),
            ("🔍 Фильтр", self.filter_),
            ("🔄 Сброс", self.unfilter),
        ]:
            btn(row1, text, cmd)

        ttk.Separator(btn_cont, orient='horizontal').pack(fill='x', pady=5)

        row2 = ttk.Frame(btn_cont)
        row2.pack(fill='x', pady=2)
        for (text, cmd) in [
            ("🔎 Поиск", self.search),
            ("📂 Открыть (2 столбца)", self.load_xy),
            ("💾 Сохранить", self.save),
            ("📁 Открыть файлы", self.load_many),
        ]:
            btn(row2, text, cmd)

        ttk.Separator(btn_cont, orient='horizontal').pack(fill='x', pady=5)

        row3 = ttk.Frame(btn_cont)
        row3.pack(fill='x', pady=2)

        self.btn_math = ttk.Button(row3, text="📐 Мат. методы", command=self._math_menu)
        self.btn_math.pack(side=tk.LEFT, padx=2, pady=2)

        self.menu = tk.Menu(self.parent, tearoff=0)
        reg_menu = tk.Menu(self.menu, tearoff=0)
        reg_menu.add_command(label="Линейная регрессия", command=self.regress)
        reg_menu.add_command(label="Полином 4-й степени", command=self.poly_reg)
        self.menu.add_cascade(label="Регрессия", menu=reg_menu)

        self.menu.add_command(label="Корреляция Пирсона", command=self.correl)

        interp_menu = tk.Menu(self.menu, tearoff=0)
        interp_menu.add_command(label="Лагранж", command=self.lagrange)
        interp_menu.add_command(label="Ньютон", command=self.newton)
        interp_menu.add_command(label="Канонический полином", command=self.canon)
        interp_menu.add_command(label="Линейный сплайн", command=self.lin_spline)
        interp_menu.add_command(label="Кубический сплайн", command=self.cub_spline)
        self.menu.add_cascade(label="Интерполяция", menu=interp_menu)

        self.menu.add_separator()
        self.menu.add_command(label="t-тест", command=self.ttest)
        self.menu.add_command(label="Удалить выбросы", command=self.outliers)
        self.menu.add_command(label="Нормализация [0,1]", command=self.norm)
        self.menu.add_command(label="Стандартизация (z)", command=self.stand)
        self.menu.add_command(label="Доверительный интервал", command=self.conf_int)

        for (text, cmd) in [
            ("📈 Корреляция", self.correl),
            ("📉 Регрессия", self.regress),
            ("🔬 t-тест", self.ttest),
            ("🧹 Выбросы", self.outliers),
        ]:
            btn(row3, text, cmd)

        main = ttk.Frame(self.parent)
        main.pack(fill='both', expand=True, padx=10, pady=10)

        left = ttk.Frame(main)
        left.pack(side=tk.LEFT, fill='both', expand=True)

        right = ttk.Frame(main, width=320)
        right.pack(side=tk.RIGHT, fill='y', padx=(10,0))
        right.pack_propagate(False)

        inp_frame = ttk.Frame(left)
        inp_frame.pack(fill='x', padx=5, pady=5)

        self.mode_btn = ttk.Button(inp_frame, text="2 столбца", command=self._toggle_mode)
        self.mode_btn.pack(side=tk.LEFT, padx=5)

        self.entry_single = ttk.Entry(inp_frame, width=20)
        self.entry_x = ttk.Entry(inp_frame, width=10)
        self.entry_y = ttk.Entry(inp_frame, width=10)

        self.entry_single.pack(side=tk.LEFT, padx=5)
        self.entry_single.bind("<Return>", lambda e: self._add_from_entry())

        self.entry_x.pack_forget()
        self.entry_y.pack_forget()
        self.entry_x.insert(0, "X")
        self.entry_y.insert(0, "Y")
        self.entry_x.bind("<Return>", lambda e: self._add_from_entry())
        self.entry_y.bind("<Return>", lambda e: self._add_from_entry())

        self.msg_lbl = ttk.Label(left, text="", foreground="green")
        self.msg_lbl.pack(pady=5)

        ttk.Label(left, text="Список (двойной клик – правка):").pack(anchor='w', padx=5)
        self.text_box = tk.Text(left, height=18, width=70, state='disabled')
        self.text_box.pack(fill='both', expand=True, padx=5, pady=5)
        self.text_box.bind("<Double-Button-1>", lambda e: self.edit())

        self.stat_box = ttk.LabelFrame(right, text="Статистика")
        self.stat_box.pack(fill='x', pady=5)
        self.stat_lbl = ttk.Label(self.stat_box, text="Нет данных", justify=tk.LEFT)
        self.stat_lbl.pack(padx=10, pady=10)

        self.info_box = ttk.LabelFrame(right, text="Информация")
        self.info_box.pack(fill='x', pady=5)
        self.info_lbl = ttk.Label(self.info_box, text="Элементов: 0", justify=tk.LEFT)
        self.info_lbl.pack(padx=10, pady=10)

        self.recent_box = ttk.LabelFrame(right, text="Последние 5")
        self.recent_box.pack(fill='x', pady=5)
        self.recent_lbl = ttk.Label(self.recent_box, text="—", justify=tk.LEFT)
        self.recent_lbl.pack(padx=10, pady=10)

    def _toggle_mode(self):
        self.two_column_mode = not self.two_column_mode
        if self.two_column_mode:
            self.mode_btn.config(text="1 столбец")
            self.entry_single.pack_forget()
            self.entry_x.pack(side=tk.LEFT, padx=5)
            self.entry_y.pack(side=tk.LEFT, padx=5)
            self.entry_x.delete(0, tk.END)
            self.entry_y.delete(0, tk.END)
            self.entry_x.insert(0, "X")
            self.entry_y.insert(0, "Y")
            self.entry_x.focus()
        else:
            self.mode_btn.config(text="2 столбца")
            self.entry_x.pack_forget()
            self.entry_y.pack_forget()
            self.entry_single.pack(side=tk.LEFT, padx=5)
            self.entry_single.delete(0, tk.END)
            self.entry_single.focus()

    def _add_from_entry(self):
        if self.two_column_mode:
            self._add_pair()
        else:
            self._add_single()

    def _add_single(self):
        s = self.entry_single.get().strip()
        if not s:
            return
        try:
            val = self._float(s)
            if self.bivar:
                self.clear()
            self.data.append(val)
            self.bivar = False
            self.x = []
            self.y = []
            self._save_hist()
            if self.f_active:
                self.unfilter()
            self.entry_single.delete(0, tk.END)
            self.show()
            self.refresh_info()
            self.msg_lbl.config(text="Добавлено", foreground="green")
        except ValueError:
            self.msg_lbl.config(text="Ошибка! Введите число", foreground="red")

    def _add_pair(self):
        sx = self.entry_x.get().strip()
        sy = self.entry_y.get().strip()
        if not sx or not sy:
            return
        try:
            x = self._float(sx)
            y = self._float(sy)
            if not self.bivar and self.data:
                self.clear()
            self.data.append((x, y))
            self.bivar = True
            self.x.append(x)
            self.y.append(y)
            self._save_hist()
            if self.f_active:
                self.unfilter()
            self.entry_x.delete(0, tk.END)
            self.entry_y.delete(0, tk.END)
            self.entry_x.insert(0, "X")
            self.entry_y.insert(0, "Y")
            self.show()
            self.refresh_info()
            self.msg_lbl.config(text="Добавлена пара", foreground="green")
        except ValueError:
            self.msg_lbl.config(text="Ошибка! Введите числа X и Y", foreground="red")

    def _float(self, s):
        s = s.strip().replace(',', '.')
        try:
            return float(s)
        except:
            raise ValueError("Не число")

    def show(self):
        arr = self.fdata if self.f_active else self.data
        self.text_box.config(state='normal')
        self.text_box.delete("1.0", tk.END)
        if self.bivar and arr:
            for x, y in arr:
                self.text_box.insert(tk.END, f"{x:.6f}\t→\t{y:.6f}\n")
        else:
            for r in arr:
                self.text_box.insert(tk.END, f"{r}\n")
        self.text_box.config(state='disabled')

    def edit(self):
        self.text_box.config(state='normal')
        try:
            first = self.text_box.index(tk.SEL_FIRST)
        except:
            tk.messagebox.showinfo("Правка", "Выделите значение")
            self.text_box.config(state='disabled')
            return
        line = int(first.split('.')[0]) - 1
        arr = self.fdata if self.f_active else self.data
        if line < 0 or line >= len(arr):
            self.text_box.config(state='disabled')
            return
        if self.bivar:
            old_x, old_y = arr[line]
            new_x = tk.simpledialog.askfloat("X", f"Текущий X = {old_x:.6f}", initialvalue=old_x)
            if new_x is not None:
                new_y = tk.simpledialog.askfloat("Y", f"Текущий Y = {old_y:.6f}", initialvalue=old_y)
                if new_y is not None:
                    if self.f_active:
                        idx = self.data.index((old_x, old_y))
                        self.data[idx] = (new_x, new_y)
                        self._save_hist()
                        self._apply_filter()
                    else:
                        self.data[line] = (new_x, new_y)
                        self._save_hist()
                    self._update_xy_from_data()
        else:
            old = arr[line]
            new = tk.simpledialog.askfloat("Изменить", f"Текущее {old}", initialvalue=old)
            if new is not None:
                if self.f_active:
                    idx = self.data.index(old)
                    self.data[idx] = new
                    self._save_hist()
                    self._apply_filter()
                else:
                    self.data[line] = new
                    self._save_hist()
        self.show()
        self.refresh_info()
        self.text_box.config(state='disabled')

    def clear(self):
        self.data.clear()
        self._save_hist()
        self.f_active = False
        self.fdata = None
        self.bivar = False
        self.x = []
        self.y = []
        self.show()
        self.update_stats()
        self.refresh_info()

    def update_stats(self):
        if self.bivar:
            self.stat_lbl.config(text="Двумерные данные\nИспользуйте\nрегрессию/корреляцию")
            return
        arr = self.fdata if self.f_active else self.data
        if not arr:
            self.stat_lbl.config(text="Нет данных")
            return
        avg = sum(arr) / len(arr)
        med = statistics.median(arr)
        mn = min(arr)
        mx = max(arr)
        std = statistics.stdev(arr) if len(arr) > 1 else 0
        self.stat_lbl.config(
            text=f"Среднее: {avg:.3f}\nМедиана: {med:.3f}\nСт. откл.: {std:.3f}\nMin: {mn}\nMax: {mx}"
        )

    def refresh_info(self):
        arr = self.fdata if self.f_active else self.data
        n = len(arr)
        self.info_lbl.config(text=f"Элементов: {n}")
        if n > 0:
            rec = arr[-5:] if n >= 5 else arr
            if self.bivar:
                txt = ", ".join([f"({x:.2f}→{y:.2f})" for (x, y) in rec])
            else:
                txt = ", ".join(str(v) for v in rec)
            self.recent_lbl.config(text=txt)
        else:
            self.recent_lbl.config(text="—")
        self.update_stats()

    def _math_menu(self):
        try:
            self.menu.post(self.btn_math.winfo_rootx(),
                           self.btn_math.winfo_rooty() + self.btn_math.winfo_height())
        except:
            pass

    # ---------- математические методы (без изменений) ----------
    def correl(self):
        if not self.bivar:
            tk.messagebox.showinfo("Корреляция", "Сначала загрузите X,Y")
            return
        r, p = sps.pearsonr(self.x, self.y)
        tk.messagebox.showinfo("Корреляция Пирсона",
                               f"r = {r:.4f}\np = {p:.4f}\n"
                               f"Интерпретация: |r|>0.7 – сильная, >0.3 – умеренная")

    def regress(self):
        if not self.bivar:
            tk.messagebox.showinfo("Регрессия", "Загрузите X,Y")
            return
        a, b, r_val, p_val, _ = sps.linregress(self.x, self.y)
        msg = f"y = {a:.4f} * x + {b:.4f}\nR² = {r_val**2:.4f}\np = {p_val:.4f}"
        tk.messagebox.showinfo("Регрессия", msg)
        plt.figure()
        plt.scatter(self.x, self.y, alpha=0.6, label='Данные')
        xr = np.array([min(self.x), max(self.x)])
        plt.plot(xr, a*xr + b, 'r-', label='Регрессия')
        plt.xlabel('X'); plt.ylabel('Y'); plt.title('Линейная регрессия')
        plt.legend(); plt.grid(True); plt.show(block=False)

    def poly_reg(self):
        if not self.bivar:
            tk.messagebox.showinfo("Полином", "Загрузите X,Y")
            return
        x = np.array(self.x)
        y = np.array(self.y)
        A = np.column_stack([np.ones_like(x), x, x**4])
        try:
            coef, _, _, _ = np.linalg.lstsq(A, y, rcond=None)
            K, D, A4 = coef
            y_pred = K + D*x + A4*x**4
            mean_y = np.mean(y)
            ss_res = np.sum((y - y_pred)**2)
            ss_tot = np.sum((y - mean_y)**2)
            R = np.sqrt(1 - ss_res/ss_tot) if ss_tot != 0 else 0
            msg = (f"Полином 4-й степени:\nZ = {K:.6f} + {D:.6f}*t + {A4:.6f}*t^4\n"
                   f"R = {R:.6f}\nСумма кв. ошибок = {ss_res:.6f}")
            tk.messagebox.showinfo("Результат", msg)
            plt.figure()
            plt.scatter(x, y, label="Данные")
            x_sm = np.linspace(min(x), max(x), 200)
            plt.plot(x_sm, K + D*x_sm + A4*x_sm**4, 'r-', label="Полином")
            plt.xlabel('t'); plt.ylabel('Z'); plt.title('Полиномиальная регрессия')
            plt.legend(); plt.grid(True); plt.show(block=False)
        except Exception as e:
            tk.messagebox.showerror("Ошибка", str(e))

    def _plot_interp(self, t, res, func, title):
        plt.figure()
        plt.scatter(self.x, self.y, color='red', label='Точки', zorder=5)
        xs = np.linspace(min(self.x), max(self.x), 200)
        plt.plot(xs, func(xs), 'b-', label=title)
        plt.plot(t, res, 'go', markersize=8, label=f't={t}')
        plt.xlabel('t'); plt.ylabel('Z'); plt.title(title)
        plt.legend(); plt.grid(True); plt.show(block=False)

    def lagrange(self):
        if not self.bivar:
            tk.messagebox.showinfo("Лагранж", "Загрузите X,Y")
            return
        t = tk.simpledialog.askfloat("Точка", "Введите t:")
        if t is None: return
        try:
            poly = spi.lagrange(self.x, self.y)
            res = poly(t)
            self._plot_interp(t, res, poly, "Интерполяция Лагранжа")
            tk.messagebox.showinfo("Результат", f"Значение в t={t}: {res:.6f}")
        except Exception as e:
            tk.messagebox.showerror("Ошибка", str(e))

    def newton(self):
        if not self.bivar:
            tk.messagebox.showinfo("Ньютон", "Загрузите X,Y")
            return
        t = tk.simpledialog.askfloat("Точка", "Введите t:")
        if t is None: return
        x = self.x
        y = self.y
        n = len(x)
        F = [[0]*n for _ in range(n)]
        for i in range(n):
            F[i][0] = y[i]
        for j in range(1, n):
            for i in range(n-j):
                F[i][j] = (F[i+1][j-1] - F[i][j-1]) / (x[i+j] - x[i])
        def pol(tt):
            res = F[0][0]
            term = 1.0
            for i in range(1, n):
                term *= (tt - x[i-1])
                res += F[0][i] * term
            return res
        res = pol(t)
        self._plot_interp(t, res, pol, "Интерполяция Ньютона")
        tk.messagebox.showinfo("Результат", f"Значение в t={t}: {res:.6f}")

    def canon(self):
        if not self.bivar:
            tk.messagebox.showinfo("Канонический", "Загрузите X,Y")
            return
        t = tk.simpledialog.askfloat("Точка", "Введите t:")
        if t is None: return
        x = self.x
        y = self.y
        n = len(x)
        A = np.vander(x, n, increasing=True)
        try:
            coef = np.linalg.solve(A, y)
            def pol(tt):
                return sum(coef[i] * (tt**i) for i in range(n))
            res = pol(t)
            self._plot_interp(t, res, pol, "Канонический полином")
            tk.messagebox.showinfo("Результат", f"Значение в t={t}: {res:.6f}")
        except np.linalg.LinAlgError:
            tk.messagebox.showerror("Ошибка", "Матрица вырождена (точки должны быть различны)")

    def lin_spline(self):
        if not self.bivar:
            tk.messagebox.showinfo("Линейный сплайн", "Загрузите X,Y")
            return
        t = tk.simpledialog.askfloat("Точка", "Введите t:")
        if t is None: return
        try:
            f = spi.interp1d(self.x, self.y, kind='linear', fill_value="extrapolate")
            res = f(t)
            self._plot_interp(t, res, f, "Линейный сплайн")
            tk.messagebox.showinfo("Результат", f"Значение в t={t}: {res:.6f}")
        except Exception as e:
            tk.messagebox.showerror("Ошибка", str(e))

    def cub_spline(self):
        if not self.bivar:
            tk.messagebox.showinfo("Кубический сплайн", "Загрузите X,Y")
            return
        t = tk.simpledialog.askfloat("Точка", "Введите t:")
        if t is None: return
        try:
            cs = spi.CubicSpline(self.x, self.y, bc_type='natural')
            res = cs(t)
            self._plot_interp(t, res, cs, "Кубический сплайн")
            tk.messagebox.showinfo("Результат", f"Значение в t={t}: {res:.6f}")
        except Exception as e:
            tk.messagebox.showerror("Ошибка", str(e))

    def ttest(self):
        if self.bivar:
            tk.messagebox.showinfo("t-тест", "Используйте одномерные данные")
            return
        if len(self.data) == 0:
            tk.messagebox.showwarning("Нет данных", "Добавьте числа")
            return
        if tk.messagebox.askyesno("t-тест", "Сравнить первую половину со второй?"):
            mid = len(self.data)//2
            g1 = self.data[:mid]
            g2 = self.data[mid:]
            if len(g1) < 2 or len(g2) < 2:
                tk.messagebox.showerror("Ошибка", "Мало данных в группах")
                return
            t_stat, p_val = sps.ttest_ind(g1, g2)
            msg = (f"Группа 1: n={len(g1)}, среднее={np.mean(g1):.3f}\n"
                   f"Группа 2: n={len(g2)}, среднее={np.mean(g2):.3f}\n"
                   f"t = {t_stat:.4f}, p = {p_val:.4f}\n"
                   f"{'Значимы' if p_val < 0.05 else 'Не значимы'}")
            tk.messagebox.showinfo("t-тест", msg)
        else:
            s = tk.simpledialog.askstring("Вторая выборка", "Введите числа через запятую")
            if s:
                try:
                    g2 = [float(x.strip().replace(',','.')) for x in s.split(',')]
                    t_stat, p_val = sps.ttest_ind(self.data, g2)
                    tk.messagebox.showinfo("t-тест", f"t = {t_stat:.4f}, p = {p_val:.4f}")
                except:
                    tk.messagebox.showerror("Ошибка", "Неверный формат")

    def outliers(self):
        if self.bivar:
            tk.messagebox.showinfo("Выбросы", "Только одномерные данные")
            return
        d = self.data[:]
        if len(d) < 4:
            tk.messagebox.showwarning("Мало данных", "Нужно минимум 4 точки")
            return
        q1 = np.percentile(d, 25)
        q3 = np.percentile(d, 75)
        iqr = q3 - q1
        lo = q1 - 1.5*iqr
        hi = q3 + 1.5*iqr
        filt = [x for x in d if lo <= x <= hi]
        rem = len(d) - len(filt)
        if tk.messagebox.askyesno("Выбросы", f"Удалено {rem} выбросов. Продолжить?"):
            self.data = filt
            self._save_hist()
            if self.f_active:
                self.unfilter()
            self.show()
            self.refresh_info()
            self.msg_lbl.config(text=f"Удалено {rem} выбросов", foreground="blue")

    def norm(self):
        if self.bivar:
            tk.messagebox.showinfo("Нормализация", "Только одномерные данные")
            return
        if not self.data: return
        mn = min(self.data)
        mx = max(self.data)
        if mx == mn:
            tk.messagebox.showwarning("Нет разброса", "Все числа одинаковы")
            return
        self.data = [(x - mn)/(mx - mn) for x in self.data]
        self._save_hist()
        self.show()
        self.refresh_info()
        self.msg_lbl.config(text="Нормализовано [0,1]")

    def stand(self):
        if self.bivar:
            tk.messagebox.showinfo("Стандартизация", "Только одномерные данные")
            return
        if len(self.data) < 2: return
        mean = np.mean(self.data)
        std = np.std(self.data)
        if std == 0:
            tk.messagebox.showwarning("Нет разброса", "σ = 0")
            return
        self.data = [(x - mean)/std for x in self.data]
        self._save_hist()
        self.show()
        self.refresh_info()
        self.msg_lbl.config(text="Стандартизировано (z-score)")

    def conf_int(self):
        if self.bivar:
            tk.messagebox.showinfo("Дов. интервал", "Только одномерные данные")
            return
        arr = self.fdata if self.f_active else self.data
        if len(arr) < 2: return
        mean = np.mean(arr)
        sem = sps.sem(arr)
        ci = sps.t.interval(0.95, len(arr)-1, loc=mean, scale=sem)
        tk.messagebox.showinfo("95% ДИ для среднего",
                               f"Среднее = {mean:.3f}\nДИ = [{ci[0]:.3f}, {ci[1]:.3f}]")

    def load_xy(self):
        path = tk.filedialog.askopenfilename(filetypes=[("CSV или TXT", "*.csv *.txt")])
        if not path: return
        try:
            with open(path, 'r', encoding='utf-8') as f:
                lines = f.readlines()
            xv, yv = [], []
            for line in lines:
                parts = line.strip().split()
                if len(parts) >= 2:
                    try:
                        xv.append(float(parts[0].replace(',','.')))
                        yv.append(float(parts[1].replace(',','.')))
                    except: pass
            if len(xv) > 1:
                self.x = xv
                self.y = yv
                self.bivar = True
                self.data = list(zip(self.x, self.y))
                self._save_hist()
                self.show()
                self.refresh_info()
                tk.messagebox.showinfo("Успех", f"Загружено {len(xv)} пар")
            else:
                tk.messagebox.showerror("Ошибка", "Файл должен содержать два числовых столбца")
        except Exception as e:
            tk.messagebox.showerror("Ошибка загрузки", str(e))

    def save(self):
        arr = self.fdata if self.f_active else self.data
        if not arr: return
        types = [("Текстовый файл", "*.txt"), ("CSV", "*.csv")]
        if XL_OK: types.append(("Excel", "*.xlsx"))
        path = tk.filedialog.asksaveasfilename(defaultextension=".txt", filetypes=types)
        if not path: return
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == '.txt':
                with open(path, 'w', encoding='utf-8') as f:
                    if self.bivar:
                        for x, y in arr: f.write(f"{x}\t{y}\n")
                    else:
                        for v in arr: f.write(f"{v}\n")
            elif ext == '.csv':
                import csv
                with open(path, 'w', encoding='utf-8', newline='') as f:
                    w = csv.writer(f)
                    if self.bivar:
                        w.writerow(["X","Y"])
                        for x, y in arr: w.writerow([x,y])
                    else:
                        w.writerow(["Value"])
                        for v in arr: w.writerow([v])
            elif ext == '.xlsx' and XL_OK:
                wb = openpyxl.Workbook()
                ws = wb.active
                if self.bivar:
                    ws.title = "XY"
                    ws.append(["X","Y"])
                    for x, y in arr: ws.append([x,y])
                else:
                    ws.title = "Data"
                    ws.append(["Value"])
                    for v in arr: ws.append([v])
                wb.save(path)
            else:
                tk.messagebox.showerror("Ошибка", "Формат не поддерживается")
                return
            tk.messagebox.showinfo("Сохранено", f"Сохранено {len(arr)} записей")
        except Exception as e:
            tk.messagebox.showerror("Ошибка сохранения", str(e))

    def _extract_nums(self, text):
        nums = []
        for t in re.findall(r'[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?', text):
            try:
                nums.append(float(t))
            except: pass
        return nums

    def load_many(self):
        if self.bivar:
            tk.messagebox.showinfo("Пакетная загрузка", "Сначала очистите данные")
            return
        paths = tk.filedialog.askopenfilenames(
            filetypes=[("Все поддерживаемые", "*.txt *.csv *.xlsx *.docx"),
                       ("Текстовые", "*.txt"), ("CSV", "*.csv"),
                       ("Excel", "*.xlsx"), ("Word", "*.docx")]
        )
        if not paths: return
        total = []
        for path in paths:
            ext = os.path.splitext(path)[1].lower()
            try:
                if ext in ('.txt', '.csv'):
                    with open(path, 'r', encoding='utf-8') as f:
                        text = f.read()
                    total += self._extract_nums(text)
                elif ext == '.xlsx' and XL_OK:
                    wb = openpyxl.load_workbook(path, data_only=True)
                    for sheet in wb.worksheets:
                        for row in sheet.iter_rows(values_only=True):
                            for cell in row:
                                if isinstance(cell, (int, float)):
                                    total.append(float(cell))
                elif ext == '.docx':
                    doc = docx.Document(path)
                    text = '\n'.join([p.text for p in doc.paragraphs])
                    total += self._extract_nums(text)
            except Exception as e:
                tk.messagebox.showerror("Ошибка", f"{os.path.basename(path)}:\n{str(e)}")
                return
        if total:
            self.data += total
            self._save_hist()
            self.show()
            self.refresh_info()
            self.msg_lbl.config(text=f"Загружено {len(total)} чисел", foreground="green")
        else:
            tk.messagebox.showwarning("Нет данных", "Не удалось извлечь числа")

    def hist_plot(self):
        if self.bivar:
            tk.messagebox.showinfo("Гистограмма", "Только одномерные данные")
            return
        arr = self.fdata if self.f_active else self.data
        if not arr: return
        plt.figure()
        plt.hist(arr, bins=10, alpha=0.7, color='blue')
        plt.title('Гистограмма')
        plt.show(block=False)

    def graph(self):
        if self.bivar:
            tk.messagebox.showinfo("График", "Только одномерные данные")
            return
        arr = self.fdata if self.f_active else self.data
        if not arr: return
        plt.figure()
        plt.plot(arr, marker='o')
        plt.title('Линейный график')
        plt.show(block=False)

    def sort(self):
        if self.bivar:
            tk.messagebox.showinfo("Сортировка", "Пары X,Y не сортируются")
            return
        arr = self.fdata if self.f_active else self.data
        if not arr: return
        arr.sort()
        if self.f_active:
            self.fdata = arr
        else:
            self.data = arr
        self._save_hist()
        self.show()

    def search(self):
        if self.bivar:
            tk.messagebox.showinfo("Поиск", "Только одномерные данные")
            return
        s = self.entry_single.get().strip() if not self.two_column_mode else ""
        if not s:
            return
        try:
            num = self._float(s)
            arr = self.fdata if self.f_active else self.data
            if num in arr:
                tk.messagebox.showinfo("Поиск", f"{num} найдено")
            else:
                tk.messagebox.showinfo("Поиск", "Не найдено")
        except:
            tk.messagebox.showerror("Ошибка", "Введите число")

    def filter_(self):
        if self.bivar:
            tk.messagebox.showinfo("Фильтр", "Только одномерные данные")
            return
        if not self.data: return
        dlg = tk.Toplevel(self.parent)
        dlg.title("Фильтр")
        dlg.geometry("300x250")
        ttk.Label(dlg, text="Условие:").pack(pady=5)
        cond_var = tk.StringVar(value=">")
        for c, d in ((">",">"), ("<","<"), (">=",">="), ("<=","<=")):
            ttk.Radiobutton(dlg, text=d, variable=cond_var, value=c).pack(anchor='w')
        ttk.Label(dlg, text="Значение:").pack()
        val_ent = ttk.Entry(dlg)
        val_ent.pack()
        def apply():
            try:
                thr = float(val_ent.get().replace(',','.'))
                cond = cond_var.get()
                self.fdata = [x for x in self.data if eval(f"{x}{cond}{thr}")]
                self.f_active = True
                self.show()
                self.refresh_info()
                dlg.destroy()
            except:
                tk.messagebox.showerror("Ошибка", "Неверное число")
        ttk.Button(dlg, text="Применить", command=apply).pack(pady=10)

    def unfilter(self):
        self.f_active = False
        self.fdata = None
        self.show()
        self.refresh_info()

    def _apply_filter(self):
        self.unfilter()


# -------------------------------------------------------------
# ВКЛАДКА ТЕКСТОВЫХ ДАННЫХ (с полными словарями)
# -------------------------------------------------------------
class TxtTab:
    def __init__(self, parent, app):
        self.parent = parent
        self.app = app
        self.lang = 'russian'
        self.langs = {'russian': 'Русский', 'english': 'English'}
        self.stops = set()
        self.stemmer = None

        # ------------------ ПОЛНЫЕ РАСШИРЕННЫЕ СЛОВАРИ (без сокращений) ------------------
        self.raw_pos = {
            'russian': [
                'хороший', 'хорошая', 'хорошее', 'хорошие', 'хорош', 'хороша', 'хорошо',
                'отличный', 'отличная', 'отличное', 'отличные', 'отличен', 'отлична', 'отлично',
                'прекрасный', 'прекрасная', 'прекрасное', 'прекрасные', 'прекрасен', 'прекрасна', 'прекрасно',
                'замечательный', 'замечательная', 'замечательное', 'замечательные', 'замечателен', 'замечательна', 'замечательно',
                'великий', 'великая', 'великое', 'великие', 'велик', 'велика', 'велико',
                'великолепный', 'великолепная', 'великолепное', 'великолепные', 'великолепен', 'великолепна', 'великолепно',
                'чудесный', 'чудесная', 'чудесное', 'чудесные', 'чудесен', 'чудесна', 'чудесно',
                'восхитительный', 'восхитительная', 'восхитительное', 'восхитительные', 'восхитителен', 'восхитительна', 'восхитительно',
                'солнечный', 'солнечная', 'солнечное', 'солнечные', 'солнечен', 'солнечна', 'солнечно',
                'успешный', 'успешная', 'успешное', 'успешные', 'успешен', 'успешна', 'успешно',
                'позитивный', 'позитивная', 'позитивное', 'позитивные', 'позитивен', 'позитивна', 'позитивно',
                'классный', 'классная', 'классное', 'классные', 'классен', 'классна', 'классно',
                'крутой', 'крутая', 'крутое', 'крутые', 'крут', 'крута', 'круто',
                'супер', 'потрясающий', 'потрясающая', 'потрясающее', 'потрясающие', 'потрясающ',
                'зашибенный', 'зашибенная', 'зашибенное', 'зашибенные', 'зашибен', 'зашибена',
                'милый', 'милая', 'милое', 'милые', 'мил', 'мила', 'мило',
                'приятный', 'приятная', 'приятное', 'приятные', 'приятен', 'приятна', 'приятно',
                'славный', 'славная', 'славное', 'славные', 'славен', 'славна', 'славно',
                'благородный', 'благородная', 'благородное', 'благородные', 'благороден', 'благородна', 'благородно',
                'лучший', 'лучшая', 'лучшее', 'лучшие', 'лучше',
                'любовь', 'любви', 'любовью', 'любовь', 'любови', 'любовей', 'любовный',
                'счастье', 'счастья', 'счастью', 'счастьем', 'счастлив', 'счастлива', 'счастливо', 'счастливый',
                'радость', 'радости', 'радостью', 'радостей', 'радостям', 'радостный',
                'добро', 'добра', 'добру', 'добром', 'добрый', 'добрая', 'доброе', 'добрые',
                'великолепие', 'великолепия', 'великолепию', 'великолепием',
                'восторг', 'восторга', 'восторгу', 'восторгом', 'восторженный',
                'улыбка', 'улыбки', 'улыбкой', 'улыбку', 'улыбаться', 'улыбнуться',
                'смех', 'смеха', 'смеху', 'смехом', 'смешной', 'смеяться',
                'праздник', 'праздника', 'празднику', 'праздником', 'праздничный',
                'победа', 'победы', 'победе', 'победой', 'победный',
                'удача', 'удачи', 'удаче', 'удачу', 'удачей', 'удачный',
                'здоровье', 'здоровья', 'здоровью', 'здоровьем', 'здоровый',
                'красота', 'красоты', 'красоте', 'красотой', 'красивый',
                'гармония', 'гармонии', 'гармонию', 'гармонией', 'гармоничный',
                'вдохновение', 'вдохновения', 'вдохновению', 'вдохновением', 'вдохновляющий',
                'успех', 'успеха', 'успеху', 'успехом', 'успехи', 'успешный',
                'любить', 'люблю', 'любишь', 'любит', 'любим', 'любите', 'любят', 'любил', 'любила', 'любили',
                'радоваться', 'радуюсь', 'радуешься', 'радуется', 'радовался', 'радовалась', 'радовались', 'радуйся',
                'восхищать', 'восхищаю', 'восхищаешь', 'восхищает', 'восхищал', 'восхищала', 'восхищаться', 'восхититься',
                'улыбаться', 'улыбаюсь', 'улыбается', 'улыбался', 'улыбалась', 'улыбнуться',
                'смеяться', 'смеюсь', 'смеется', 'смеялся', 'смеялась',
                'побеждать', 'побеждаю', 'побеждает', 'побеждал', 'победить',
                'радовать', 'радую', 'радует', 'радовал', 'радовало',
                'хорошо', 'отлично', 'прекрасно', 'замечательно', 'великолепно', 'чудесно', 'классно', 'круто',
                'позитивно', 'радостно', 'счастливо', 'доброжелательно', 'успешно', 'благополучно', 'превосходно',
                'великодушно', 'гармонично', 'вдохновенно', 'победно', 'удачно', 'красиво', 'мило', 'приятно'
            ],
            'english': [
                'good', 'better', 'best', 'well', 'goodly', 'goodness',
                'great', 'greater', 'greatest', 'greatly',
                'excellent', 'excellently', 'excellence',
                'wonderful', 'wonderfully', 'wonderfulness',
                'fantastic', 'fantastically',
                'amazing', 'amazingly', 'amazed', 'amazingness',
                'beautiful', 'beautifully', 'beauty',
                'nice', 'nicer', 'nicest', 'nicely', 'niceness',
                'pleased', 'pleasing', 'pleasantly', 'pleasure',
                'glorious', 'gloriously', 'glory',
                'superb', 'superbly', 'super',
                'brilliant', 'brilliantly', 'brilliance',
                'awesome', 'awesomely', 'awesomeness',
                'perfect', 'perfectly', 'perfection',
                'lovely', 'lovelier', 'loveliest', 'loveliness',
                'delightful', 'delightfully', 'delight',
                'splendid', 'splendidly', 'splendor',
                'positive', 'positively', 'positivity',
                'favorable', 'favorably', 'favor',
                'enjoyable', 'enjoyably', 'enjoyment',
                'cheerful', 'cheerfully', 'cheer', 'cheerfulness',
                'ecstatic', 'ecstatically', 'ecstasy',
                'happy', 'happier', 'happiest', 'happily', 'happiness',
                'joy', 'joyful', 'joyfully', 'joyous', 'joyousness',
                'love', 'loves', 'loving', 'loved', 'lovely', 'lover',
                'victory', 'victories', 'victorious', 'victoriously',
                'smile', 'smiles', 'smiling', 'smiled',
                'laughter', 'laugh', 'laughs', 'laughing', 'laughed', 'laughable',
                'celebration', 'celebrate', 'celebrating', 'celebrated', 'celebratory',
                'success', 'successes', 'successful', 'successfully',
                'luck', 'lucky', 'luckier', 'luckiest', 'luckily', 'good luck',
                'health', 'healthy', 'healthier', 'healthiest', 'healthily',
                'harmony', 'harmonious', 'harmoniously',
                'inspiration', 'inspirational', 'inspiring', 'inspired',
                'enjoy', 'enjoys', 'enjoyed', 'enjoying',
                'celebrate', 'celebrates', 'celebrated', 'celebrating',
                'succeed', 'succeeds', 'succeeded', 'succeeding'
            ]
        }

        self.raw_neg = {
            'russian': [
                'плохой', 'плохая', 'плохое', 'плохие', 'плох', 'плоха', 'плохо',
                'ужасный', 'ужасная', 'ужасное', 'ужасные', 'ужасен', 'ужасна', 'ужасно',
                'отвратительный', 'отвратительная', 'отвратительное', 'отвратительные', 'отвратителен', 'отвратительна', 'отвратительно',
                'грустный', 'грустная', 'грустное', 'грустные', 'грустен', 'грустна', 'грустно',
                'мерзкий', 'мерзкая', 'мерзкое', 'мерзкие', 'мерзок', 'мерзка', 'мерзко',
                'скверный', 'скверная', 'скверное', 'скверные', 'скверен', 'скверна', 'скверно',
                'негативный', 'негативная', 'негативное', 'негативные', 'негативен', 'негативна', 'негативно',
                'дурной', 'дурная', 'дурное', 'дурные', 'дурен', 'дурна', 'дурно',
                'кошмарный', 'кошмарная', 'кошмарное', 'кошмарные', 'кошмарен', 'кошмарна', 'кошмарно',
                'гадкий', 'гадкая', 'гадкое', 'гадкие', 'гадок', 'гадка', 'гадко',
                'злой', 'злая', 'злое', 'злые', 'зол', 'зла', 'зло',
                'жестокий', 'жестокая', 'жестокое', 'жестокие', 'жесток', 'жестока', 'жестоко',
                'противный', 'противная', 'противное', 'противные', 'противен', 'противна', 'противно',
                'безнадежный', 'безнадежная', 'безнадежное', 'безнадежные', 'безнадежен', 'безнадежна', 'безнадежно',
                'ненависть', 'ненависти', 'ненавистью', 'ненавистный',
                'зло', 'зла', 'злу', 'злом', 'злые',
                'печаль', 'печали', 'печалью', 'печальный', 'печально',
                'беда', 'беды', 'беде', 'беду', 'бедой', 'бедный',
                'проблема', 'проблемы', 'проблеме', 'проблему', 'проблемой', 'проблемный',
                'боль', 'боли', 'болью', 'болевой', 'больно',
                'страх', 'страха', 'страху', 'страхом', 'страхи', 'страшный', 'страшно',
                'жестокость', 'жестокости', 'жестокостью',
                'мерзость', 'мерзости', 'мерзостью',
                'гадость', 'гадости', 'гадостью',
                'слеза', 'слезы', 'слезу', 'слезой', 'слёзный',
                'плач', 'плача', 'плачу', 'плачем', 'плакать',
                'потеря', 'потери', 'потерю', 'потерей', 'потерянный',
                'болезнь', 'болезни', 'болезнью', 'больной',
                'несчастье', 'несчастья', 'несчастью', 'несчастьем', 'несчастный',
                'разочарование', 'разочарования', 'разочарованию', 'разочарованием', 'разочарованный',
                'агрессия', 'агрессии', 'агрессию', 'агрессией', 'агрессивный',
                'конфликт', 'конфликта', 'конфликту', 'конфликтом', 'конфликтный',
                'тоска', 'тоски', 'тоске', 'тоской', 'тоскливый', 'тоскливо',
                'уныние', 'уныния', 'унынию', 'унынием', 'унылый',
                'ненавидеть', 'ненавижу', 'ненавидишь', 'ненавидит', 'ненавидел', 'ненавидела',
                'злиться', 'злюсь', 'злишься', 'злится', 'злился', 'злилась',
                'плакать', 'плачу', 'плачешь', 'плачет', 'плакал', 'плакала', 'плачь',
                'болеть', 'болею', 'болеешь', 'болеет', 'болел', 'болела',
                'страдать', 'страдаю', 'страдаешь', 'страдает', 'страдал', 'страдала', 'страдающий',
                'унывать', 'унываю', 'унывает', 'унывал', 'унывала',
                'разочаровывать', 'разочаровываю', 'разочаровывает', 'разочаровывал',
                'конфликтовать', 'конфликтую', 'конфликтует', 'конфликтовал',
                'плохо', 'ужасно', 'отвратительно', 'грустно', 'скверно', 'негативно', 'дурно', 'кошмарно',
                'больно', 'страшно', 'жестоко', 'гадко', 'мерзко', 'тоскливо', 'безнадежно', 'уныло', 'печально'
            ],
            'english': [
                'bad', 'worse', 'worst', 'badly', 'badness',
                'terrible', 'terribly', 'terribleness',
                'awful', 'awfully', 'awfulness',
                'horrible', 'horribly', 'horrid',
                'sad', 'sadder', 'saddest', 'sadly', 'sadness',
                'nasty', 'nastier', 'nastiest', 'nastily', 'nastiness',
                'negative', 'negatively', 'negativity',
                'poor', 'poorer', 'poorest', 'poorly', 'poverty',
                'disgusting', 'disgustingly', 'disgust',
                'miserable', 'miserably', 'misery',
                'unpleasant', 'unpleasantly', 'unpleasantness',
                'evil', 'evilly', 'evilness',
                'dreadful', 'dreadfully',
                'lousy', 'lousily',
                'rotten', 'rottenly',
                'unhappy', 'unhappier', 'unhappiest', 'unhappily', 'unhappiness',
                'gloomy', 'gloomily', 'gloom',
                'hopeless', 'hopelessly', 'hopelessness',
                'painful', 'painfully',
                'fearful', 'fearfully',
                'violent', 'violently',
                'aggressive', 'aggressively',
                'hate', 'hatred', 'hateful', 'hater',
                'anger', 'angry', 'angrier', 'angriest', 'angrily',
                'pain', 'pains', 'painful', 'painfully',
                'sorrow', 'sorrows', 'sorrowful', 'sorrowfully',
                'grief', 'griefs', 'grievous', 'grieve',
                'fear', 'fears', 'fearful', 'fearfully',
                'trouble', 'troubles', 'troublesome',
                'disaster', 'disasters', 'disastrous', 'disastrously',
                'failure', 'failures', 'failing',
                'loss', 'losses', 'lost',
                'illness', 'illnesses', 'ill', 'sickness',
                'suffering', 'sufferings', 'suffer',
                'disappointment', 'disappointments', 'disappointed',
                'frustration', 'frustrations', 'frustrated',
                'violence', 'violent', 'violently',
                'cry', 'cries', 'crying', 'cried',
                'suffer', 'suffers', 'suffered', 'suffering',
                'hurt', 'hurts', 'hurting', 'hurted',
                'harm', 'harms', 'harmed', 'harming',
                'kill', 'kills', 'killed', 'killing',
                'disappoint', 'disappoints', 'disappointed', 'disappointing',
                'frustrate', 'frustrates', 'frustrated', 'frustrating'
            ]
        }

        self._build()
        self._set_lang()

    def _set_lang(self):
        try:
            self.stops = set(nltk.corpus.stopwords.words(self.lang))
        except:
            self.stops = set()
        try:
            self.stemmer = nltk.stem.SnowballStemmer(self.lang)
        except:
            self.stemmer = None
        self.pos_stems = set()
        self.neg_stems = set()
        if self.stemmer:
            for w in self.raw_pos.get(self.lang, []):
                self.pos_stems.add(self.stemmer.stem(w))
            for w in self.raw_neg.get(self.lang, []):
                self.neg_stems.add(self.stemmer.stem(w))
        self.refresh_stats()

    def _chg_lang(self, event=None):
        self.lang = self.lang_var.get()
        self._set_lang()
        self.sentiment()

    def _build(self):
        btn_frame = ttk.Frame(self.parent)
        btn_frame.pack(pady=10, fill='x')

        def btn(text, cmd):
            b = ttk.Button(btn_frame, text=text, command=cmd)
            b.pack(side=tk.LEFT, padx=2, pady=2)
            return b

        for (text, cmd) in [
            ("📂 Загрузить", self.load),
            ("💾 Сохранить", self.save),
            ("🗑️ Очистить", self.clear),
            ("🔍 Поиск", self.search),
            ("☁️ Облако слов", self.cloud),
            ("🚫 Удалить стоп-слова", self.del_stops),
            ("🌱 Стемминг", self.stem_words),
        ]:
            btn(text, cmd)

        main = ttk.Frame(self.parent)
        main.pack(fill='both', expand=True, padx=10, pady=10)

        left = ttk.Frame(main)
        left.pack(side=tk.LEFT, fill='both', expand=True)

        right = ttk.Frame(main, width=350)
        right.pack(side=tk.RIGHT, fill='y', padx=(10,0))
        right.pack_propagate(False)

        self.text_box = tk.Text(left, height=25, font=("Consolas", 10),
                                undo=True, autoseparators=True, maxundo=100)
        self.text_box.pack(fill='both', expand=True)
        self.text_box.bind("<KeyRelease>", lambda e: self.refresh_stats())
        self.text_box.edit_reset()

        self.stat_box = ttk.LabelFrame(right, text="Статистика")
        self.stat_box.pack(fill='x', pady=5)
        self.stat_lbl = ttk.Label(self.stat_box, text="Нет данных", justify=tk.LEFT)
        self.stat_lbl.pack(padx=10, pady=10)

        self.sent_box = ttk.LabelFrame(right, text="Тональность")
        self.sent_box.pack(fill='x', pady=5)
        self.sent_lbl = ttk.Label(self.sent_box, text="—", justify=tk.LEFT)
        self.sent_lbl.pack(padx=10, pady=10)

        self.lang_box = ttk.LabelFrame(right, text="Язык текста")
        self.lang_box.pack(fill='x', pady=5)
        self.lang_var = tk.StringVar(value='russian')
        self.lang_cb = ttk.Combobox(self.lang_box, textvariable=self.lang_var,
                                    values=list(self.langs.keys()), state='readonly')
        self.lang_cb.pack(padx=10, pady=10, fill='x')
        self.lang_cb.bind('<<ComboboxSelected>>', self._chg_lang)

    def undo(self, event=None):
        try:
            self.text_box.edit_undo()
            self.refresh_stats()
        except:
            pass

    def redo(self, event=None):
        try:
            self.text_box.edit_redo()
            self.refresh_stats()
        except:
            pass

    def get_txt(self):
        return self.text_box.get("1.0", tk.END).strip()

    def words_from(self, text):
        return re.findall(r'\b\w+\b', text.lower(), flags=re.UNICODE)

    def refresh_stats(self):
        text = self.get_txt()
        if not text:
            self.stat_lbl.config(text="Нет текста")
            self.sent_lbl.config(text="—")
            return
        chars = len(text)
        words = self.words_from(text)
        wc = len(words)
        sents = text.count('.') + text.count('!') + text.count('?')
        self.stat_lbl.config(text=f"Символов: {chars}\nСлов: {wc}\nПредложений: {sents}")
        self.sentiment()

    def load(self):
        path = tk.filedialog.askopenfilename(
            filetypes=[("Все поддерживаемые", "*.txt *.csv *.pdf *.docx"),
                       ("Текстовые", "*.txt"), ("CSV", "*.csv"),
                       ("PDF", "*.pdf"), ("Word", "*.docx")]
        )
        if not path: return
        content = ""
        try:
            if path.lower().endswith(('.txt', '.csv')):
                with open(path, 'r', encoding='utf-8') as f:
                    content = f.read()
            elif path.lower().endswith('.pdf'):
                with open(path, 'rb') as f:
                    reader = pypdf.PdfReader(f)
                    for page in reader.pages:
                        content += page.extract_text()
            elif path.lower().endswith('.docx'):
                doc = docx.Document(path)
                content = '\n'.join([p.text for p in doc.paragraphs])
            else:
                tk.messagebox.showerror("Ошибка", "Неподдерживаемый формат")
                return
        except Exception as e:
            tk.messagebox.showerror("Ошибка загрузки", str(e))
            return
        if content:
            self.text_box.delete("1.0", tk.END)
            self.text_box.insert("1.0", content)
            self.text_box.edit_reset()
            self.refresh_stats()
        else:
            tk.messagebox.showwarning("Пустой файл", "Не удалось извлечь текст")

    def save(self):
        text = self.get_txt()
        if not text:
            tk.messagebox.showwarning("Нет данных", "Нечего сохранять")
            return
        types = [("Текстовый", "*.txt"), ("CSV", "*.csv")]
        if XL_OK: types.append(("Excel", "*.xlsx"))
        path = tk.filedialog.asksaveasfilename(defaultextension=".txt", filetypes=types)
        if not path: return
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == '.txt':
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(text)
            elif ext == '.csv':
                import csv
                with open(path, 'w', encoding='utf-8', newline='') as f:
                    writer = csv.writer(f)
                    writer.writerow(["Text"])
                    writer.writerow([text])
            elif ext == '.xlsx' and XL_OK:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Text"
                ws.append(["Content"])
                for line in text.split('\n'):
                    ws.append([line])
                wb.save(path)
            else:
                tk.messagebox.showerror("Ошибка", "Неподдерживаемый формат")
                return
            tk.messagebox.showinfo("Сохранено", "Текст сохранён")
        except Exception as e:
            tk.messagebox.showerror("Ошибка сохранения", str(e))

    def clear(self):
        self.text_box.delete("1.0", tk.END)
        self.text_box.edit_reset()
        self.refresh_stats()

    def search(self):
        word = tk.simpledialog.askstring("Поиск", "Слово/фрагмент:")
        if word:
            count = self.get_txt().lower().count(word.lower())
            tk.messagebox.showinfo("Поиск", f"Вхождений: {count}")

    def cloud(self):
        text = self.get_txt()
        if not text: return
        wc_obj = wc.WordCloud(width=800, height=400,
                              background_color='white',
                              stopwords=self.stops).generate(text)
        plt.figure()
        plt.imshow(wc_obj, interpolation='bilinear')
        plt.axis('off')
        plt.show(block=False)

    def sentiment(self):
        text = self.get_txt()
        if not text:
            self.sent_lbl.config(text="Нет текста")
            return
        if not self.stemmer:
            self.sent_lbl.config(text="Стеммер недоступен")
            return
        words = self.words_from(text)
        if not words:
            self.sent_lbl.config(text="Нет слов")
            return
        stems = [self.stemmer.stem(w) for w in words]
        pos_count = sum(1 for s in stems if s in self.pos_stems)
        neg_count = sum(1 for s in stems if s in self.neg_stems)
        total = pos_count + neg_count
        if total == 0:
            sentiment_text = "Нейтрально"
            detail = "Позитивных/негативных слов не найдено"
        else:
            pos_ratio = pos_count / total
            if pos_ratio > 0.6:
                sentiment_text = f"Позитивное ({pos_ratio*100:.0f}%)"
            elif pos_ratio < 0.4:
                sentiment_text = f"Негативное ({(1-pos_ratio)*100:.0f}%)"
            else:
                sentiment_text = "Нейтрально"
            detail = f"Позитив: {pos_count}, Негатив: {neg_count}"
        self.sent_lbl.config(text=f"{sentiment_text}\n{detail}")

    def del_stops(self):
        text = self.get_txt()
        words = self.words_from(text)
        filtered = [w for w in words if w not in self.stops]
        self.text_box.delete("1.0", tk.END)
        self.text_box.insert("1.0", ' '.join(filtered))
        self.text_box.edit_reset()
        self.refresh_stats()

    def stem_words(self):
        if not self.stemmer:
            tk.messagebox.showwarning("Стемминг", "Для выбранного языка стемминг недоступен")
            return
        text = self.get_txt()
        words = self.words_from(text)
        stemmed = [self.stemmer.stem(w) for w in words]
        self.text_box.delete("1.0", tk.END)
        self.text_box.insert("1.0", ' '.join(stemmed))
        self.text_box.edit_reset()
        self.refresh_stats()


if __name__ == "__main__":
    root = tk.Tk()
    style = ttk.Style()
    style.theme_use('clam')
    default_font = ('Segoe UI Emoji', 10)
    style.configure('TButton', font=default_font, padding=6)
    style.configure('TLabel', font=default_font)
    style.configure('TLabelframe.Label', font=default_font)
    style.configure('TEntry', font=default_font)
    app = App(root)
    root.mainloop()
